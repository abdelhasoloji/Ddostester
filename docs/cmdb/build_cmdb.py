#!/usr/bin/env python3
"""Build the WEDOS ShieldTest project CMDB workbook. No secret VALUES are stored."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/ddostester/docs/cmdb/shieldtest-cmdb.xlsx"
BRAND = "003484"      # WEDOS navy
BRAND2 = "0A4AA0"
LIGHT = "EAF0F8"
WARN = "FFF3CD"
FONT = "Arial"

hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=BRAND)
title_font = Font(name=FONT, bold=True, color=BRAND, size=16)
sub_font = Font(name=FONT, italic=True, color="555555", size=10)
base_font = Font(name=FONT, size=10)
bold_font = Font(name=FONT, bold=True, size=10)
warn_fill = PatternFill("solid", fgColor=WARN)
zebra = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
top = Alignment(vertical="top", wrap_text=True)

wb = Workbook()


def sheet(title, headers, rows, widths, intro=None, notes=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    r = 1
    ws.cell(r, 1, f"WEDOS ShieldTest — {title}").font = title_font
    r += 1
    if intro:
        ws.cell(r, 1, intro).font = sub_font
        r += 1
    r += 1
    head_row = r
    for c, h in enumerate(headers, 1):
        cell = ws.cell(head_row, c, h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    r += 1
    for i, row in enumerate(rows):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = base_font
            cell.alignment = top
            cell.border = border
            if i % 2 == 1:
                cell.fill = zebra
        r += 1
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    if notes:
        r += 1
        for n in notes:
            cell = ws.cell(r, 1, n)
            cell.font = sub_font
            r += 1
    return ws, head_row


# ── README ───────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 100
ws["A1"] = "WEDOS ShieldTest — Project CMDB"
ws["A1"].font = title_font
rows = [
    ("Purpose", "Single source of truth for everything that must be preserved for the ShieldTest engagement: hosts, network, repos, endpoints, env, and a secrets REGISTRY (names/locations only — never values)."),
    ("Project", "WEDOS ShieldTest — DDoS-resilience control plane. Engagement: rehearse the zeroBS L7 playbook (P26-20371) against the authorized target bank.cz and tune WEDOS mitigation rules."),
    ("Security boundary", "Bounded, allowlisted, authorized, capped HTTP only. No exploit / amplification / spoofing / evasion. Class-C volumetric on the isolated lab runner only. Never hit ENISA production."),
    ("5-phase model", "Recon -> Plan -> Execute -> Measure -> Report (bounded analogues for phases 3-4). See docs/five-phase-model.md."),
    ("SECRETS RULE", "This workbook stores NO secret values. The 'Secrets Registry' sheet lists what exists and where it is kept (password manager). Rotate the WEDOS root passwords after first login."),
    ("Owner", "abdelha@gmail.com (WEDOS, internal usage)"),
    ("Last updated", "2026-09-09"),
    ("Tabs", "VPS Inventory | Network & Firewall | Roles & Services | Repositories & CICD | Endpoints | Environment Variables | Secrets Registry | Engagement | Change Log"),
]
r = 3
for k, v in rows:
    a = ws.cell(r, 1, k)
    a.font = bold_font
    a.alignment = top
    b = ws.cell(r, 2, v)
    b.font = base_font
    b.alignment = top
    if k in ("SECRETS RULE", "Security boundary"):
        a.fill = warn_fill
        b.fill = warn_fill
    r += 1

# ── VPS Inventory ────────────────────────────────────────────────────────────
vps_ws, vps_hdr = sheet(
    "VPS Inventory",
    ["Hostname", "Order #", "Role", "Public IPv4", "CIDR", "Gateway v4",
     "IPv6 primary (host part masked)", "vCPU", "RAM", "Disk", "OS",
     "KVM/VNC (host:port)", "Established", "Expiration", "Status", "Notes"],
    [
        ["vm32047", "3326007110", "Control plane (VPS 1)", "31.31.74.228", "/24", "31.31.74.1",
         "2a02:2b88:2:1::7d2f:x /64 (gw ::1)", "4", "16 GB", "120 GB SSD", "Debian 13",
         "32047.vm47.wedos.net:42047", "2026-09-09 08:21", "2027-09-09", "actif",
         "App + UI + load engines + local Postgres"],
        ["vm32046", "3326007109", "Lab runner (VPS 2)", "31.31.74.187", "/24", "31.31.74.1",
         "2a02:2b88:2:1::7d2e:x /64 (gw ::1)", "2", "8 GB", "60 GB SSD", "Debian 13",
         "32046.vm46.wedos.net:42046", "2026-09-09 08:21", "2027-09-09", "actif",
         "Isolated class-C bench runner; egress to VPS1 API only"],
        ["vm32044", "3326007107", "Spare / standby", "46.28.109.92", "/24", "46.28.109.1",
         "2a02:2b88:2:1::7d2c:x /64 (gw ::1)", "2", "8 GB", "60 GB SSD", "Debian 13",
         "32044.vm44.wedos.net:42044", "2026-09-09 08:19", "2027-09-09", "actif", "Reserve"],
        ["vm32045", "3326007108", "Spare / standby", "31.31.73.165", "/24", "31.31.73.1",
         "2a02:2b88:2:1::7d2d:x /64 (gw ::1)", "1", "4 GB", "30 GB SSD", "Debian 13",
         "32045.vm45.wedos.net:42045", "2026-09-09 08:19", "2027-09-09", "actif", "Reserve"],
    ],
    [10, 12, 20, 14, 6, 13, 30, 6, 8, 12, 10, 26, 17, 12, 8, 34],
    intro="4 x WEDOS VPS SSD (Debian 13). Provider: WEDOS Internet a.s. Internal usage (payment cancelled). IPv6 host part is masked in the client panel — read the exact value at client.wedos.com.",
    notes=[
        "DNS (v4): 46.28.108.2, 31.31.72.3   |   DNS (v6): 2a02:2b88:2:1::2552:1, 2a02:2b88:2:1::af4:1",
        "VPS 1 and VPS 2 share the 31.31.74.0/24 subnet -> the runner reaches VPS1's API on the same /24.",
        "Active VPS: 4 of 4 (vm32044, vm32045, vm32046, vm32047)",
    ],
)

# ── Network & Firewall ───────────────────────────────────────────────────────
sheet(
    "Network & Firewall",
    ["Host", "Direction", "Port/Proto", "Source / Destination", "Purpose", "Rule (ufw)"],
    [
        ["VPS 1 (31.31.74.228)", "inbound", "22/tcp", "admin IP", "SSH admin", "ufw allow 22/tcp"],
        ["VPS 1", "inbound", "8080/tcp", "31.31.74.187 (VPS 2)", "runner claims jobs", "ufw allow from 31.31.74.187 to any port 8080 proto tcp"],
        ["VPS 1", "inbound", "8080/tcp", "admin IP", "console/API over tunnel", "ufw allow from <ADMIN_IP> to any port 8080 proto tcp"],
        ["VPS 1", "loopback", "5432/tcp", "127.0.0.1", "local Postgres (not exposed)", "n/a (bind loopback)"],
        ["VPS 1", "outbound", "443/tcp", "bank.cz", "authorized test target", "default allow out"],
        ["VPS 2 (31.31.74.187)", "inbound", "22/tcp", "admin IP", "SSH admin", "ufw allow 22/tcp"],
        ["VPS 2", "outbound", "8080/tcp", "31.31.74.228 (VPS 1)", "post verdicts to API", "ufw allow out to 31.31.74.228 port 8080 proto tcp"],
        ["VPS 2", "outbound", "53", "DNS resolvers", "name resolution", "ufw allow out 53"],
        ["VPS 2", "default", "all", "-", "fail-closed", "ufw default deny incoming + outgoing"],
    ],
    [22, 10, 12, 24, 26, 52],
    intro="Firewall posture. VPS 2 is fail-closed: deny all, allow ONLY egress to VPS1 API + DNS, and admin SSH. Replace <ADMIN_IP> with your management address.",
)

# ── Roles & Services ─────────────────────────────────────────────────────────
sheet(
    "Roles & Services",
    ["Service", "Host", "Port", "Unit / Process", "Config", "Notes"],
    [
        ["Control-plane API + UI", "VPS 1", "8080", "systemd: ddostester (deploy/setup.sh)", "/etc/ddostester.env", "Fastify + React; health at /health"],
        ["PostgreSQL", "VPS 1", "5432", "postgresql", "local role 'shieldtest'", "loopback only; DATABASE_URL"],
        ["Load engines", "VPS 1", "-", "in-process (builtin) + k6/vegeta if installed", "-", "bounded HTTP generators"],
        ["Class-C lab runner", "VPS 2", "-", "systemd: shieldtest-lab-runner", "BENCH_API_URL, BENCH_RUNNER_TOKEN", "claims governed jobs, posts read-only verdicts"],
        ["Deploy runner", "ac02 (self-hosted)", "-", "GitHub Actions self-hosted runner", "-", "runs Deploy workflow via workflow_run"],
    ],
    [26, 18, 8, 40, 30, 40],
)

# ── Repositories & CI/CD ─────────────────────────────────────────────────────
sheet(
    "Repositories & CICD",
    ["Repository", "Visibility", "Purpose", "Dev branch", "Flow", "Notes"],
    [
        ["abdelhasoloji/ddostester", "private", "Control-plane app (this project)", "claude/nouveau-projet-git-xi95n7", "branch -> ff main -> CI + E2E -> Deploy (ac02)", "CI gate includes test:coverage (backend 73/frontend 80 floors)"],
        ["abdelhasoloji/shieldtest-lab", "private", "Isolated class-C lab runner", "-", "cloned on VPS 2 by host-lab.sh", "npm run runner"],
        ["abdelhasoloji/lexkyc", "private", "Separate KYC product (not this project)", "-", "-", "Different repo; listed to avoid confusion"],
    ],
    [30, 10, 34, 30, 40, 44],
    notes=[
        "Deploy keys (read-only) on ddostester: root@VPS1 (bootstrap fetch) + shieldtest@VPS1 (service clone).",
        "Bootstrap scripts: deploy/provision/host-app.sh (VPS1), deploy/provision/host-lab.sh (VPS2).",
    ],
)

# ── Endpoints ────────────────────────────────────────────────────────────────
sheet(
    "Endpoints",
    ["Name", "URL / Address", "Auth", "Notes"],
    [
        ["Control-plane API/UI", "http://31.31.74.228:8080", "session (AUTH_USERS) / read token", "switch to https only if TLS fronted"],
        ["Health check", "http://31.31.74.228:8080/health", "open", "storage, allowedTargets, protection, hardCaps"],
        ["Recon (prep phase)", "GET http://31.31.74.228:8080/api/recon?target=bank.cz", "campaign:read", "passive + published-surface recon"],
        ["KVM/VNC (per host)", "see VPS Inventory (SSL VNC, e.g. SSVNC)", "KVM password (WEDOS email)", "out-of-band console"],
        ["WEDOS client panel", "https://client.wedos.com/", "WEDOS account", "power ops, ISO mount, IP/reverse records"],
        ["WEDOS KB (VPS)", "https://kb.wedos.com/en/kategorie/virtual-servers-vps-ssd/", "open", "vendor documentation"],
    ],
    [24, 56, 26, 40],
)

# ── Environment Variables ────────────────────────────────────────────────────
sheet(
    "Environment Variables",
    ["Host", "Variable", "Example / value shape", "Required", "Purpose"],
    [
        ["VPS 1", "NODE_ENV", "production", "yes", "runtime mode"],
        ["VPS 1", "PORT", "8080", "yes", "API/UI listen port"],
        ["VPS 1", "DATABASE_URL", "postgresql://shieldtest:<pw>@127.0.0.1:5432/shieldtest", "yes", "DB connection (pw = secret)"],
        ["VPS 1", "AUTH_JWT_SECRET", "<32-byte hex>", "yes", "session signing (secret)"],
        ["VPS 1", "AUTH_USERS", '[{"email":..,"role":..,"passwordHash":..}]', "yes", "operator accounts"],
        ["VPS 1", "ALLOWED_TARGETS", "bank.cz", "yes", "target allowlist (authorization)"],
        ["VPS 1", "LAB_TARGETS", "lab", "yes", "class-C governance label (NOT a connect target)"],
        ["VPS 1", "BENCH_RUNNER_TOKEN", "<24-byte hex>", "yes", "app<->lab bearer (secret; same on VPS2)"],
        ["VPS 1", "PROTECTION_METRICS_URL", "<scrubber metrics endpoint>", "recommended", "measured TTD/TTM (else estimated)"],
        ["VPS 1", "PROTECTION_SOURCE_LABEL", "wedos-scrubbing", "optional", "protection source label"],
        ["VPS 1", "PROTECTION_TIMEOUT_MS", "5000", "optional", "protection hook timeout"],
        ["VPS 2", "BENCH_API_URL", "http://31.31.74.228:8080", "yes", "control-plane API base"],
        ["VPS 2", "BENCH_RUNNER_TOKEN", "<same as VPS1>", "yes", "bearer (secret; must match VPS1)"],
    ],
    [8, 24, 46, 14, 40],
    intro="Env keys read by host-app.sh / host-lab.sh. Values shown are SHAPES ONLY. Real secret values live in the password manager (see Secrets Registry).",
)

# ── Secrets Registry (NO VALUES) ─────────────────────────────────────────────
sheet(
    "Secrets Registry",
    ["Secret", "Purpose", "Scope / where used", "Stored in", "Generated", "Rotation status"],
    [
        ["DB_PASSWORD", "Postgres app-role password", "VPS 1 (/etc/ddostester.env, DATABASE_URL)", "Password manager", "openssl rand -hex 24", "set at provisioning"],
        ["AUTH_JWT_SECRET", "App session signing", "VPS 1 env", "Password manager", "openssl rand -hex 32", "set at provisioning"],
        ["BENCH_RUNNER_TOKEN", "app<->lab runner bearer", "VPS 1 + VPS 2 (identical)", "Password manager", "openssl rand -hex 24", "set at provisioning"],
        ["SSH root password (per VPS)", "initial root login", "each VPS", "WEDOS setup e-mail (transient)", "WEDOS", "ROTATE after first login -> key-only auth"],
        ["Deploy key: root@VPS1", "git clone ddostester (bootstrap)", "GitHub repo Deploy keys", "GitHub + VPS1 /root/.ssh", "ssh-keygen ed25519", "read-only"],
        ["Deploy key: shieldtest@VPS1", "git clone ddostester (service)", "GitHub repo Deploy keys", "GitHub + VPS1 shieldtest home", "ssh-keygen ed25519 (by host-app.sh)", "read-only"],
        ["KVM/VNC password (per VPS)", "out-of-band console", "each VPS", "WEDOS setup e-mail", "WEDOS", "n/a"],
        ["PROTECTION_METRICS_URL", "scrubber hook endpoint", "VPS 1 env", "Password manager / infra note", "WEDOS scrubber", "PENDING — not yet wired"],
    ],
    [26, 28, 34, 26, 26, 34],
    intro="REGISTRY ONLY — no secret values are stored here. If a value ever needs sharing, use the password manager, never this file or the repo.",
    notes=[
        "Reminder: the three generated secrets were pasted into a chat transcript on 2026-09-09; rotate if that transcript is considered exposed.",
        "After first SSH login on each VPS: change root password, install an admin SSH key, disable password auth.",
    ],
)

# ── Engagement ───────────────────────────────────────────────────────────────
sheet(
    "Engagement",
    ["Item", "Value"],
    [
        ["Target", "bank.cz (authorized external/internal test target)"],
        ["Allowlist", "ALLOWED_TARGETS=bank.cz"],
        ["Protection under test", "WEDOS (scrubbing / edge)"],
        ["Reference playbook", "zeroBS Playbook P26-20371 (ENISA/CRA portal, DRS v3, L7)"],
        ["ShieldTest playbook", "pretest-cra-l7 — 12 bounded L7 steps, DRS 3->5 (coverage 12/14)"],
        ["Out of model (attested)", "HPACK bomb, WAF-cracker, credential brute-force, L3/4 handshake (bench)"],
        ["DRS target cap", "4.5 (customer-defined targets capped)"],
        ["Loop", "Recon -> Plan (gap ladder) -> Execute -> Measure (WRS) -> Report -> tune-and-replay"],
        ["Key docs", "docs/five-phase-model.md, docs/pretest-p26-20371.md, docs/tune-and-replay.md, docs/provisioning-4vps.md"],
    ],
    [26, 86],
)

# ── Change Log ───────────────────────────────────────────────────────────────
sheet(
    "Change Log",
    ["Date", "Change", "By"],
    [
        ["2026-09-09", "4 x WEDOS VPS provisioned (vm32044-47), all active", "abdelha / WEDOS"],
        ["2026-09-09", "Roles mapped: vm32047=control plane, vm32046=lab runner; vm32044/45 spare", "Claude"],
        ["2026-09-09", "Shared secrets generated (DB_PASSWORD, AUTH_JWT_SECRET, BENCH_RUNNER_TOKEN)", "abdelha"],
        ["2026-09-09", "CMDB created", "Claude"],
        ["2026-09-09", "GOTCHA: SSH/console to VPS filtered while WEDOS VPN is connected (VPN filters the 31.31.74.0/24 subnet). Disconnect VPN for admin access.", "abdelha"],
        ["2026-09-09", "VPS 1 control plane LIVE: host-app.sh bootstrap, systemd ddostester, storage=postgres, ALLOWED_TARGETS=bank.cz, auth on (eng@wedos/officer@wedos). /health OK.", "Claude/abdelha"],
        ["2026-09-09", "First cycle validated end-to-end vs bank.cz (Recon->Plan->Execute->Measure->Report). Smoke test GET-Flood DRS3 = WRS 97 A+, author!=approver enforced.", "abdelha"],
        ["2026-09-09", "Deploy-script bugs fixed in repo: setup.sh (pipefail && / persist BENCH_RUNNER_TOKEN+LAB_TARGETS) and host-app.sh (sudo prereq, shieldtest real home).", "Claude"],
        ["", "PENDING: VPS 2 lab runner (host-lab.sh) + scrubber hook PROTECTION_METRICS_URL + full pretest-cra-l7 suite.", ""],
    ],
    [14, 74, 20],
    intro="Append one row per material change (host added, IP change, secret rotation, role change).",
)

wb.save(OUT)
print("saved", OUT)
